import re
import unicodedata
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from .scoring import build_ats_result


WORD_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MISSION_EMAIL = "boubacar@osito-solution.com"


def _safe_filename(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = "".join(character for character in normalized if not unicodedata.combining(character))
    return re.sub(r"[^A-Za-z0-9]+", "_", ascii_value).strip("_")[:80] or "Mission"


def targeted_cv_filename(opportunity, extension: str = "docx") -> str:
    label = _safe_filename(opportunity.title)
    company = _safe_filename(opportunity.company) if opportunity.company else ""
    suffix = f"_{company}" if company else ""
    return f"Boubacar_DIABY_{label}{suffix}.{extension}"


def _replace_text_runs(paragraph, value: str) -> None:
    """Replace visible text while preserving drawing-only runs such as the CV photo."""
    text_runs = [run for run in paragraph.runs if run.text]
    if not text_runs:
        paragraph.add_run(value)
        return
    text_runs[0].text = value
    for run in text_runs[1:]:
        run.text = ""


def _headline(ats: dict) -> str:
    keywords = []
    for keyword in ats["matched_keywords"]:
        label = keyword.upper() if keyword in {"api", "crm", "cpq", "lwc", "soql"} else keyword.title()
        if label not in keywords:
            keywords.append(label)
    return " | ".join(keywords[:4]) or "Salesforce | Architecture CRM | Delivery"


def _set_run_font(run, font_name: str) -> None:
    """Set every Word font slot so the result is stable across renderers."""
    run.font.name = font_name
    run_properties = run._element.get_or_add_rPr()
    run_fonts = run_properties.get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        run_fonts.set(qn(f"w:{slot}"), font_name)


def _harmonize_impact_business_fonts(document) -> None:
    """Keep the Impact Business bullets in the CV's main Arial typeface."""
    in_impact_section = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.upper() == "IMPACT BUSINESS":
            in_impact_section = True
            continue
        if in_impact_section and text.upper() == "EXPÉRIENCES PROFESSIONNELLES":
            break
        if in_impact_section:
            for run in paragraph.runs:
                _set_run_font(run, "Arial")


def _remove_trailing_empty_paragraphs(document) -> None:
    """Avoid blank trailing pages caused by empty paragraphs in the master CV."""
    while document.paragraphs and not document.paragraphs[-1].text.strip():
        paragraph_element = document.paragraphs[-1]._element
        paragraph_element.getparent().remove(paragraph_element)


def build_targeted_cv(master_content: bytes, opportunity, profile) -> bytes:
    """Create a tailored copy of the master CV without changing the stored original."""
    document = Document(BytesIO(master_content))
    # Internal assessments (for example AssessFirst notes) may inform the
    # private ATS analysis, but must never be copied verbatim into a CV.
    ats = build_ats_result(opportunity, profile, include_soft_skill_profile=False)
    visible_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    if not visible_paragraphs:
        raise ValueError("Le CV maître ne contient aucun paragraphe exploitable")

    title = opportunity.title.strip() or ats["suggested_title"]
    _replace_text_runs(visible_paragraphs[0], f"{title}\n{_headline(ats)}")

    for paragraph in visible_paragraphs[:10]:
        for run in paragraph.runs:
            if "@" in run.text:
                run.text = re.sub(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", MISSION_EMAIL, run.text)

    summary_paragraph = next(
        (
            paragraph
            for paragraph in visible_paragraphs[1:12]
            if len(paragraph.text.strip()) >= 120 and paragraph.text.strip() != paragraph.text.strip().upper()
        ),
        None,
    )
    if summary_paragraph is not None:
        _replace_text_runs(summary_paragraph, ats["tailored_summary"])

    _harmonize_impact_business_fonts(document)
    _remove_trailing_empty_paragraphs(document)

    document.core_properties.title = f"CV ciblé — {title}"
    document.core_properties.subject = opportunity.company or "Mission Salesforce"
    document.core_properties.keywords = ", ".join(ats["matched_keywords"])

    output = BytesIO()
    document.save(output)
    return output.getvalue()
