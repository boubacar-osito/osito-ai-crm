from io import BytesIO
from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn

from app.cv_generation import build_targeted_cv, targeted_cv_filename


def test_build_targeted_cv_updates_title_and_summary_without_changing_master():
    source = Document()
    source.add_paragraph("ARCHITECTE SOLUTION SALESFORCE\nCRM | DELIVERY")
    source.add_paragraph("Boubacar DIABY")
    source.add_paragraph("boubacar.diaby@example.com")
    source.add_paragraph(
        "Architecte CRM et Salesforce senior avec une expérience complète du cadrage, "
        "de l'architecture et du pilotage de programmes de transformation complexes auprès de grands comptes."
    )
    buffer = BytesIO()
    source.save(buffer)
    master = buffer.getvalue()
    opportunity = SimpleNamespace(
        title="Chef de projet Salesforce",
        company="Groupe Média",
        description="Pilotage agile d'un projet Salesforce, coordination du delivery et des API.",
    )
    profile = SimpleNamespace(
        title="Architecte CRM/Salesforce senior",
        summary="Plus de 20 ans d'expérience SI et plus de 10 ans sur Salesforce.",
        skills=["Salesforce", "CRM", "Architecture", "Agile", "API", "Delivery"],
        cv_text="Salesforce CRM architecture agile API delivery",
        soft_skill_profile="Leadership, pédagogie et sens du collectif.",
    )

    generated = build_targeted_cv(master, opportunity, profile)
    result = Document(BytesIO(generated))

    assert master == buffer.getvalue()
    assert "Chef de projet Salesforce" in result.paragraphs[0].text
    assert "Salesforce" in result.paragraphs[0].text
    assert "Plus de 20 ans" in result.paragraphs[3].text
    assert result.paragraphs[2].text == "boubacar@osito-solution.com"
    assert result.core_properties.title == "CV ciblé — Chef de projet Salesforce"


def test_build_targeted_cv_never_exports_internal_soft_skill_assessment():
    source = Document()
    source.add_paragraph("Architecte Salesforce")
    source.add_paragraph("Boubacar DIABY")
    source.add_paragraph("boubacar@example.com")
    source.add_paragraph(
        "Architecte CRM senior avec une solide expérience de l'architecture, "
        "du cadrage et du delivery de programmes Salesforce internationaux."
    )
    buffer = BytesIO()
    source.save(buffer)
    opportunity = SimpleNamespace(
        title="Chef de Projet Salesforce - Media - Paris",
        company="EASY PARTNER",
        description="Pilotage Salesforce, delivery et intégration CRM",
    )
    profile = SimpleNamespace(
        title="Architecte CRM-SI senior",
        summary="Plus de 20 ans d'expérience SI, dont plus de 10 ans sur Salesforce.",
        skills=["Salesforce", "CRM", "Delivery", "Intégration"],
        cv_text="Salesforce CRM delivery intégration",
        soft_skill_profile="Synthèse AssessFirst SWIPE / DRIVE / BRAIN strictement interne",
    )

    generated = build_targeted_cv(buffer.getvalue(), opportunity, profile)
    generated_text = "\n".join(paragraph.text for paragraph in Document(BytesIO(generated)).paragraphs)

    assert "AssessFirst" not in generated_text
    assert "SWIPE" not in generated_text
    assert "Posture professionnelle" not in generated_text


def test_build_targeted_cv_harmonizes_impact_business_fonts():
    source = Document()
    source.add_paragraph("Architecte Salesforce")
    source.add_paragraph("Boubacar DIABY")
    source.add_paragraph("boubacar@example.com")
    source.add_paragraph(
        "Architecte CRM senior avec une solide expérience de l'architecture, "
        "du cadrage et du delivery de programmes Salesforce internationaux."
    )
    source.add_paragraph("IMPACT BUSINESS")
    source.add_paragraph("+40 % résolution des incidents")
    mixed_font = source.add_paragraph().add_run("+20 organisations Salesforce")
    mixed_font.font.name = "Cambria"
    source.add_paragraph("EXPÉRIENCES PROFESSIONNELLES")
    untouched = source.add_paragraph().add_run("Responsable CRM")
    untouched.font.name = "Cambria"
    buffer = BytesIO()
    source.save(buffer)
    opportunity = SimpleNamespace(
        title="Chef de Projet Salesforce",
        company="Groupe Média",
        description="Pilotage Salesforce et delivery CRM",
    )
    profile = SimpleNamespace(
        title="Architecte CRM-SI senior",
        summary="Plus de 20 ans d'expérience SI, dont plus de 10 ans sur Salesforce.",
        skills=["Salesforce", "CRM", "Delivery"],
        cv_text="Salesforce CRM delivery",
        soft_skill_profile="",
    )

    generated = Document(BytesIO(build_targeted_cv(buffer.getvalue(), opportunity, profile)))
    impact_runs = generated.paragraphs[5].runs + generated.paragraphs[6].runs

    assert all(run.font.name == "Arial" for run in impact_runs)
    assert all(run._element.rPr.rFonts.get(qn("w:hAnsi")) == "Arial" for run in impact_runs)
    assert generated.paragraphs[8].runs[0].font.name == "Cambria"


def test_targeted_cv_filename_is_safe_and_identifies_the_mission():
    opportunity = SimpleNamespace(title="Architecte Salesforce – Île-de-France", company="Société Média")

    assert targeted_cv_filename(opportunity) == "Boubacar_DIABY_Architecte_Salesforce_Ile_de_France_Societe_Media.docx"
    assert targeted_cv_filename(opportunity, "pdf").endswith(".pdf")
